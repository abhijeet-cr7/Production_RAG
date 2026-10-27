"""DOCX guard: reject macros, embedded objects, remote templates and zip bombs.

A .docx is a zip of XML parts, so it can be attacked at the container level
(decompression bombs), the XML level (XXE) and the Office level (VBA macros,
OLE/ActiveX objects, remote templates fetched on open). Everything here runs
inside the parser sandbox.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from offline_pipeline.safety.errors import UnsafeContent
from offline_pipeline.safety.xml_guard import safe_fromstring
from offline_pipeline.tables.blocks import (
    HEADING,
    MAX_TABLE_CELLS,
    PARAGRAPH,
    TABLE,
    Block,
    Table,
    blocks_to_json,
    blocks_to_text,
    infer_header_rows,
    trim_grid,
)
from offline_pipeline.tables.numbers import detect_units

MAX_ENTRIES = 5000
MAX_TOTAL_UNCOMPRESSED = 200 * 1024 * 1024
MAX_PART_BYTES = 50 * 1024 * 1024
MAX_RATIO = 100  # compressed→uncompressed, applied to parts over 1 MB

_FORBIDDEN_NAMES = ("vbaproject.bin", "vbadata.xml")
_FORBIDDEN_PREFIXES = ("word/embeddings/", "word/activex/")
_FORBIDDEN_REL_TYPES = ("attachedtemplate", "/oleobject", "/frame", "/subdocument", "/package")
_REQUIRED = "[Content_Types].xml"


def inspect_docx(path: Path) -> None:
    """Raise :class:`UnsafeContent` unless the file is a plain, bounded .docx."""
    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise UnsafeContent("docx_not_zip", "File is not a valid DOCX.") from exc

    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ENTRIES:
            raise UnsafeContent("docx_too_many_entries", "DOCX has too many parts.")
        names = {i.filename for i in infos}
        if _REQUIRED not in names or "word/document.xml" not in names:
            raise UnsafeContent("docx_malformed", "File is not a valid DOCX.")

        total = 0
        for info in infos:
            lowered = info.filename.lower()
            if lowered.endswith(_FORBIDDEN_NAMES) or lowered.startswith(_FORBIDDEN_PREFIXES):
                raise UnsafeContent("docx_active_content", "DOCX contains macros or embedded objects.")
            if info.filename.startswith("/") or ".." in info.filename.split("/"):
                raise UnsafeContent("docx_path_traversal", "DOCX contains an unsafe part name.")
            if info.flag_bits & 0x1:
                raise UnsafeContent("docx_encrypted", "Encrypted DOCX files are not accepted.")
            if info.file_size > MAX_PART_BYTES:
                raise UnsafeContent("docx_part_too_large", "A DOCX part is too large.")
            if info.file_size > 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > MAX_RATIO:
                raise UnsafeContent("docx_bomb", "DOCX compression ratio is suspicious.")
            total += info.file_size
            if total > MAX_TOTAL_UNCOMPRESSED:
                raise UnsafeContent("docx_bomb", "DOCX expands to too much data.")

        for info in infos:
            if info.filename.endswith((".xml", ".rels")):
                data = _read_bounded(archive, info)
                root = safe_fromstring(data)
                if info.filename.endswith(".rels"):
                    _check_relationships(root)


def _read_bounded(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> bytes:
    # file_size is attacker-declared, so read with a hard ceiling instead.
    with archive.open(info) as handle:
        data = handle.read(MAX_PART_BYTES + 1)
    if len(data) > MAX_PART_BYTES:
        raise UnsafeContent("docx_part_too_large", "A DOCX part is too large.")
    return data


def _check_relationships(root: Any) -> None:
    for rel in root:
        rel_type = (rel.get("Type") or "").lower()
        external = (rel.get("TargetMode") or "").lower() == "external"
        if any(rel_type.endswith(t) for t in _FORBIDDEN_REL_TYPES):
            raise UnsafeContent("docx_active_content", "DOCX references an embedded or remote object.")
        if external and rel_type.endswith("/template"):
            raise UnsafeContent("docx_remote_template", "DOCX loads a remote template.")


def extract_docx(path: Path, *, max_chars: int) -> dict[str, Any]:
    """Return visible paragraphs, headings and tables of a (previously inspected) DOCX, in order."""
    from docx import Document
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    inspect_docx(path)
    document = Document(str(path))
    blocks: list[Block] = []
    total = 0
    hidden = 0
    cells_seen = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = _visible_text(paragraph)
            hidden += _hidden_run_count(paragraph)
            if not text.strip():
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = _heading_level(style)
            blocks.append(Block(kind=HEADING if level else PARAGRAPH, text=text, level=level))
            total += len(text) + 1
        elif tag == "tbl":
            grid, seen, table_hidden = _table_grid(DocxTable(child, document))
            cells_seen += seen
            hidden += table_hidden
            if cells_seen > MAX_TABLE_CELLS:
                raise UnsafeContent("table_too_large", "Document tables have too many cells.")
            grid = trim_grid(grid)
            if len(grid) >= 2:
                header_rows, body = infer_header_rows(grid)
                scale, currency = detect_units(" ".join(" ".join(r) for r in header_rows[:2]))
                blocks.append(Block(kind=TABLE, table=Table(
                    header_rows=header_rows, rows=body, scale=scale, currency=currency)))
                total += sum(len(c) for row in grid for c in row)
        if total > max_chars:
            raise UnsafeContent("text_too_large", "Extracted text is too large.")

    # A short paragraph just above a table is its caption.
    for i, block in enumerate(blocks[:-1]):
        nxt = blocks[i + 1]
        if (block.kind == PARAGRAPH and nxt.kind == TABLE and nxt.table is not None
                and not nxt.table.caption and len(block.text) < 160):
            nxt.table.caption = block.text
            scale, currency = detect_units(block.text)
            nxt.table.scale = nxt.table.scale or scale
            nxt.table.currency = nxt.table.currency or currency
            block.text = ""
    blocks = [b for b in blocks if b.kind != PARAGRAPH or b.text]
    return {
        "text": blocks_to_text(blocks),
        "metadata": {"file_type": "docx", "hidden_runs_removed": hidden},
        "blocks": blocks_to_json(blocks),
    }


def _heading_level(style_name: str) -> int:
    name = style_name.lower()
    if name == "title":
        return 1
    if name.startswith("heading"):
        digits = "".join(ch for ch in name if ch.isdigit())
        return min(6, max(1, int(digits))) if digits else 1
    return 0


def _table_grid(table: Any) -> tuple[list[list[str]], int, int]:
    """Rows of visible cell text; a merged cell appears once, not once per column it spans."""
    grid: list[list[str]] = []
    count = 0
    hidden = 0
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            count += 1
            if cell._tc is previous:   # horizontally merged: same underlying cell
                cells.append("")
                continue
            previous = cell._tc
            texts = []
            for paragraph in cell.paragraphs:
                texts.append(_visible_text(paragraph))
                hidden += _hidden_run_count(paragraph)
            cells.append(" ".join(t for t in texts if t))
        grid.append(cells)
    return grid, count, hidden


def _run_is_hidden(run: Any) -> bool:
    if run.font.hidden:  # w:vanish
        return True
    size = run.font.size
    return size is not None and size.pt <= 2


def _visible_text(paragraph: Any) -> str:
    return "".join(run.text for run in paragraph.runs if not _run_is_hidden(run))


def _hidden_run_count(paragraph: Any) -> int:
    return sum(1 for run in paragraph.runs if _run_is_hidden(run) and run.text.strip())
